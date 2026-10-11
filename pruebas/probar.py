#!/usr/bin/env python3
"""Ejecuta las baterías de escenarios contra el servicio de decisiones y genera un informe.

Uso:
  python3 pruebas/probar.py                          # todas las baterías de ./escenarios contra http://localhost:8080
  python3 pruebas/probar.py --url http://host:8080 --carpeta escenarios --salida informe
  python3 pruebas/probar.py escenarios/FRIOCAMPO_planificar_entrega.json   # una batería concreta

Genera en la carpeta de salida:
  informe.html      informe visual (verde = coincide, rojo = no coincide)
  resultados.xlsx   los mismos resultados en Excel (si está instalado openpyxl)
  resultados.json   resultados en bruto
Si existe la variable GITHUB_STEP_SUMMARY (GitHub Actions), escribe además un resumen en la ejecución.
Devuelve código 1 si algún escenario no coincide con lo esperado.

Solo usa la biblioteca estándar de Python (openpyxl es opcional).
"""
import argparse, datetime, glob, html, json, os, sys, urllib.error, urllib.parse, urllib.request

TOL = 0.005


# ------------------------------------------------------------------ comparación
def obtener(resp, ruta):
    if ruta in ('resultado', '$'):
        return resp, True
    actual = resp
    for parte in ruta.split('.'):
        if isinstance(actual, dict) and parte in actual:
            actual = actual[parte]
        else:
            return None, False
    return actual, True


def iguales(esperado, obtenido):
    if isinstance(esperado, bool) or isinstance(obtenido, bool):
        return type(esperado) == type(obtenido) and esperado == obtenido
    if isinstance(esperado, (int, float)) and isinstance(obtenido, (int, float)):
        return abs(esperado - obtenido) <= TOL
    if isinstance(esperado, list) and isinstance(obtenido, list):
        return len(esperado) == len(obtenido) and all(iguales(a, b) for a, b in zip(esperado, obtenido))
    if isinstance(esperado, dict) and isinstance(obtenido, dict):
        return all(k in obtenido and iguales(v, obtenido[k]) for k, v in esperado.items())
    return esperado == obtenido


def llamar(url, cuerpo):
    datos = json.dumps(cuerpo).encode('utf-8')
    req = urllib.request.Request(url, data=datos, method='POST', headers={'Content-Type': 'application/json', 'Accept': 'application/json'})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return r.status, r.read().decode('utf-8')
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode('utf-8', 'replace')
    except Exception as e:  # servicio caído, URL mal escrita…
        return 0, str(e)


def mensajes_error(texto):
    try:
        d = json.loads(texto)
    except Exception:
        return [texto[:300]] if texto else []
    if not isinstance(d, dict): return []
    ms = d.get('messages', [])   # primero los errores, después los avisos
    return [m.get('message', '') for m in ms if m.get('severity') == 'ERROR'] + [m.get('message', '') for m in ms if m.get('severity') != 'ERROR']


def ruta_servicio(base, bateria):
    partes = [bateria['modelo']] + ([bateria['servicio']] if bateria.get('servicio') else [])
    return base.rstrip('/') + '/' + '/'.join(urllib.parse.quote(p) for p in partes)


def ejecutar_bateria(base, bateria):
    url = ruta_servicio(base, bateria)
    filas = []
    for esc in bateria.get('escenarios', []):
        http_esperado = esc.get('http', 200)
        codigo, texto = llamar(url, esc.get('entrada', {}))
        try:
            resp = json.loads(texto)
        except Exception:
            resp = texto
        comprobaciones = []
        if codigo != http_esperado:
            comprobaciones.append({'campo': 'código HTTP', 'esperado': http_esperado, 'obtenido': codigo, 'ok': False})
        msgs = mensajes_error(texto) if codigo != 200 else []
        if 'mensaje' in (esc.get('esperado') or {}):
            buscado = esc['esperado']['mensaje']
            hallado = next((m for m in msgs if buscado.lower() in m.lower()), None)
            comprobaciones.append({'campo': 'mensaje de error', 'esperado': buscado, 'obtenido': hallado or (msgs[0] if msgs else '(sin mensaje)'), 'ok': hallado is not None})
        if codigo == 200:
            for campo, esperado in (esc.get('esperado') or {}).items():
                if campo == 'mensaje': continue
                obtenido, existe = obtener(resp, campo)
                comprobaciones.append({'campo': campo, 'esperado': esperado, 'obtenido': obtenido if existe else '(no aparece)',
                                       'ok': existe and iguales(esperado, obtenido)})
        if not esc.get('esperado') and codigo == http_esperado:
            estado = 'SIN ESPERADO'
        elif codigo == 0:
            estado = 'ERROR'
        else:
            estado = 'OK' if all(c['ok'] for c in comprobaciones) else 'FALLO'
        filas.append({'id': esc.get('id', ''), 'tipo': esc.get('tipo', ''), 'descripcion': esc.get('descripcion', ''),
                      'estado': estado, 'http': codigo, 'comprobaciones': comprobaciones, 'entrada': esc.get('entrada', {}),
                      'respuesta': resp, 'mensajes': msgs})
    return {'titulo': bateria.get('titulo', bateria.get('fichero', '')), 'fichero': bateria.get('fichero', ''), 'url': url,
            'descripcion': bateria.get('descripcion', ''), 'escenarios': filas}


# ------------------------------------------------------------------ informes
def fmt(v):
    if isinstance(v, str):
        return v
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return json.dumps(v, ensure_ascii=False)


CSS = """
:root{--azul:#0098CD;--azul-o:#0A5F94;--ok:#1E7B34;--ok-f:#E6F4EA;--ko:#B3261E;--ko-f:#FCE8E6;--gris:#5F6368;--borde:#DADCE0;--fondo:#F8F9FA}
*{box-sizing:border-box}body{font-family:Calibri,'Segoe UI',Arial,sans-serif;margin:0;background:var(--fondo);color:#202124}
header{background:var(--azul-o);color:#fff;padding:18px 24px}header h1{margin:0;font-size:22px}header p{margin:4px 0 0;opacity:.85;font-size:14px}
main{padding:20px 24px;max-width:1300px;margin:auto}
.kpis{display:flex;gap:12px;flex-wrap:wrap;margin-bottom:18px}.kpi{background:#fff;border:1px solid var(--borde);border-radius:8px;padding:10px 16px;min-width:130px}
.kpi b{display:block;font-size:26px}.kpi.ok b{color:var(--ok)}.kpi.ko b{color:var(--ko)}
section{background:#fff;border:1px solid var(--borde);border-radius:8px;margin-bottom:18px;overflow:hidden}
section h2{margin:0;padding:12px 16px;font-size:17px;background:#EAF6FB;border-bottom:1px solid var(--borde)}
section .meta{padding:6px 16px;color:var(--gris);font-size:13px;border-bottom:1px solid var(--borde)}
table{border-collapse:collapse;width:100%;font-size:13.5px}th{background:var(--azul);color:#fff;text-align:left;padding:7px 8px;font-weight:600}
td{padding:7px 8px;border-bottom:1px solid var(--borde);vertical-align:top}
.st{font-weight:700;border-radius:4px;padding:2px 8px;display:inline-block;white-space:nowrap}
.OK{background:var(--ok-f);color:var(--ok)}.FALLO,.ERROR{background:var(--ko-f);color:var(--ko)}.SIN{background:#FEF7E0;color:#8A6100}
.chk{margin:0;padding:0;list-style:none}.chk li{margin:1px 0}.chk .bad{color:var(--ko);font-weight:600}.chk .good{color:var(--ok)}
details summary{cursor:pointer;color:var(--azul-o)}pre{background:#F1F3F4;padding:8px;border-radius:4px;max-height:260px;overflow:auto;font-size:12px;white-space:pre-wrap}
code{font-family:Consolas,monospace}footer{color:var(--gris);font-size:12px;text-align:center;padding:12px}
"""


def informe_html(resultados, base, fecha):
    total = sum(len(b['escenarios']) for b in resultados)
    ok = sum(1 for b in resultados for e in b['escenarios'] if e['estado'] == 'OK')
    ko = sum(1 for b in resultados for e in b['escenarios'] if e['estado'] in ('FALLO', 'ERROR'))
    h = ['<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">',
         '<title>Pruebas de decisiones</title><style>%s</style></head><body>' % CSS,
         '<header><h1>Pruebas de los servicios de decisión</h1><p>Servicio: <code>%s</code> · %s</p></header><main>' % (html.escape(base), fecha),
         '<div class="kpis"><div class="kpi"><b>%d</b>escenarios</div><div class="kpi ok"><b>%d</b>coinciden</div><div class="kpi ko"><b>%d</b>no coinciden</div><div class="kpi"><b>%d</b>baterías</div></div>' % (total, ok, ko, len(resultados))]
    for b in resultados:
        n_ok = sum(1 for e in b['escenarios'] if e['estado'] == 'OK')
        h.append('<section><h2>%s · %d/%d</h2><div class="meta">%s<br><code>POST %s</code></div>' % (
            html.escape(b['titulo']), n_ok, len(b['escenarios']), html.escape(b['descripcion']), html.escape(b['url'])))
        h.append('<table><tr><th>ID</th><th>Tipo</th><th>Escenario</th><th>Resultado</th><th>Comprobaciones (esperado → obtenido)</th><th>Detalle</th></tr>')
        for e in b['escenarios']:
            chk = ''.join('<li class="%s">%s %s: %s → %s</li>' % ('good' if c['ok'] else 'bad', '✔' if c['ok'] else '✘', html.escape(c['campo']),
                                                                  html.escape(fmt(c['esperado'])), html.escape(fmt(c['obtenido']))) for c in e['comprobaciones'])
            if e['mensajes']:
                chk += ''.join('<li class="bad">%s</li>' % html.escape(m[:200]) for m in e['mensajes'][:3])
            est = e['estado']
            h.append('<tr><td>%s</td><td>%s</td><td>%s</td><td><span class="st %s">%s</span><br><small>HTTP %s</small></td><td><ul class="chk">%s</ul></td>'
                     '<td><details><summary>petición y respuesta</summary><pre>%s</pre><pre>%s</pre></details></td></tr>' % (
                         html.escape(e['id']), html.escape(e['tipo']), html.escape(e['descripcion']), est.split()[0], est, e['http'], chk,
                         html.escape(json.dumps(e['entrada'], ensure_ascii=False, indent=1)), html.escape(json.dumps(e['respuesta'], ensure_ascii=False, indent=1))))
        h.append('</table></section>')
    h.append('</main><footer>Generado por pruebas/probar.py · Sistemas de Ayuda a la Decisión (UNIR)</footer></body></html>')
    return '\n'.join(h)


def informe_xlsx(resultados, ruta):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return False
    wb = Workbook(); ws = wb.active; ws.title = 'Resultados'
    cab = ['Batería', 'ID', 'Tipo', 'Escenario', 'Resultado', 'HTTP', 'Campo', 'Esperado', 'Obtenido', '¿Coincide?']
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='0098CD')
    verde, rojo = PatternFill('solid', fgColor='E6F4EA'), PatternFill('solid', fgColor='FCE8E6')
    for b in resultados:
        for e in b['escenarios']:
            filas = e['comprobaciones'] or [{'campo': '', 'esperado': '', 'obtenido': '', 'ok': e['estado'] == 'OK'}]
            for c in filas:
                ws.append([b['titulo'], e['id'], e['tipo'], e['descripcion'], e['estado'], e['http'], c['campo'], fmt(c['esperado']), fmt(c['obtenido']), 'Sí' if c['ok'] else 'No'])
                ws.cell(ws.max_row, 5).fill = verde if e['estado'] == 'OK' else rojo
                ws.cell(ws.max_row, 10).fill = verde if c['ok'] else rojo
    for col, w in zip('ABCDEFGHIJ', [34, 6, 16, 52, 11, 7, 26, 34, 34, 11]):
        ws.column_dimensions[col].width = w
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    wb.save(ruta)
    return True


def resumen_md(resultados, base):
    total = sum(len(b['escenarios']) for b in resultados)
    ok = sum(1 for b in resultados for e in b['escenarios'] if e['estado'] == 'OK')
    lineas = ['## Pruebas de los servicios de decisión', '', '**%d de %d escenarios coinciden con lo esperado.**' % (ok, total), '']
    for b in resultados:
        n_ok = sum(1 for e in b['escenarios'] if e['estado'] == 'OK')
        lineas += ['### %s (%d/%d)' % (b['titulo'], n_ok, len(b['escenarios'])), '', '| | ID | Tipo | Escenario | Detalle |', '|---|---|---|---|---|']
        for e in b['escenarios']:
            icono = '✅' if e['estado'] == 'OK' else ('⚠️' if e['estado'] == 'SIN ESPERADO' else '❌')
            mal = '; '.join('%s: esperado %s, obtenido %s' % (c['campo'], fmt(c['esperado']), fmt(c['obtenido'])) for c in e['comprobaciones'] if not c['ok'])
            lineas.append('| %s | %s | %s | %s | %s |' % (icono, e['id'], e['tipo'], e['descripcion'], mal.replace('|', '/') or ''))
        lineas.append('')
    return '\n'.join(lineas)


# ------------------------------------------------------------------ principal
def main():
    ap = argparse.ArgumentParser(description='Ejecuta baterías de escenarios contra un servicio de decisiones DMN.')
    ap.add_argument('baterias', nargs='*', help='ficheros de batería (por defecto, todos los de --carpeta)')
    ap.add_argument('--url', default=os.environ.get('SERVICIO_URL', 'http://localhost:8080'))
    ap.add_argument('--carpeta', default='escenarios')
    ap.add_argument('--salida', default='informe')
    a = ap.parse_args()

    ficheros = a.baterias or sorted(glob.glob(os.path.join(a.carpeta, '*.json')))
    if not ficheros:
        print('No hay baterías en', a.carpeta); return 0
    resultados = []
    for f in ficheros:
        with open(f, encoding='utf-8') as fh:
            bat = json.load(fh)
        bat['fichero'] = os.path.basename(f)
        r = ejecutar_bateria(a.url, bat)
        resultados.append(r)
        n_ok = sum(1 for e in r['escenarios'] if e['estado'] == 'OK')
        print('%-50s %2d/%2d' % (r['titulo'], n_ok, len(r['escenarios'])))
        for e in r['escenarios']:
            if e['estado'] != 'OK':
                print('   %-5s %-12s %s' % (e['id'], e['estado'], e['descripcion']))

    os.makedirs(a.salida, exist_ok=True)
    fecha = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
    with open(os.path.join(a.salida, 'informe.html'), 'w', encoding='utf-8') as fh:
        fh.write(informe_html(resultados, a.url, fecha))
    with open(os.path.join(a.salida, 'resultados.json'), 'w', encoding='utf-8') as fh:
        json.dump(resultados, fh, ensure_ascii=False, indent=1)
    xlsx = informe_xlsx(resultados, os.path.join(a.salida, 'resultados.xlsx'))
    print('Informe:', os.path.join(a.salida, 'informe.html') + (' y resultados.xlsx' if xlsx else ''))
    if os.environ.get('GITHUB_STEP_SUMMARY'):
        with open(os.environ['GITHUB_STEP_SUMMARY'], 'a', encoding='utf-8') as fh:
            fh.write(resumen_md(resultados, a.url))
    return 1 if any(e['estado'] in ('FALLO', 'ERROR') for b in resultados for e in b['escenarios']) else 0


if __name__ == '__main__':
    sys.exit(main())
